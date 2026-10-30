"""Excel export: one sheet of transactions plus a summary sheet."""

from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .models import ParseResult

HEADERS = ["#", "Date", "Narration", "Ref / UTR", "Debit", "Credit", "Balance", "Balance check", "Ledger", "Tally group", "Voucher type", "Note", "Flags", "Review"]
FLAG_FILL = PatternFill("solid", fgColor="FFF4CE")
FAIL_FILL = PatternFill("solid", fgColor="FDE2E1")


def build_excel(result: ParseResult) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Transactions"
    ws.append(HEADERS)
    for c in ws[1]:
        c.font = Font(bold=True)
    for t in result.transactions:
        ws.append([
            t.row_no,
            t.txn_date,
            t.narration,
            t.instrument_no or t.ref,
            float(t.debit) if t.debit else None,
            float(t.credit) if t.credit else None,
            float(t.balance) if t.balance is not None else None,
            "OK" if t.balance_ok else ("MISMATCH" if t.balance_ok is False else "n/a"),
            t.ledger,
            t.group,
            t.voucher_type,
            t.note,
            ", ".join(t.flags),
            {"approved": "Approved", "excluded": "Left out of Tally"}.get(t.review_state, ""),
        ])
        row = ws[ws.max_row]
        row[1].number_format = "DD-MM-YYYY"
        for idx in (4, 5, 6):
            row[idx].number_format = "#,##,##0.00"
        if t.balance_ok is False:
            for c in row:
                c.fill = FAIL_FILL
        elif "needs_review" in t.flags:
            for c in row:
                c.fill = FLAG_FILL
    widths = [5, 12, 60, 20, 14, 14, 16, 14, 28, 20, 12, 30, 30, 18]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    for row in ws.iter_rows(min_row=2, min_col=3, max_col=3):
        row[0].alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"

    s = wb.create_sheet("Summary")
    for k, v in result.summary().items():
        s.append([k.replace("_", " ").capitalize(), ", ".join(v) if isinstance(v, list) else v])
    s.column_dimensions["A"].width = 22
    s.column_dimensions["B"].width = 60

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
