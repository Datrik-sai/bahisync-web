from datetime import date
from decimal import Decimal
from xml.etree import ElementTree as ET

import httpx
from openpyxl import load_workbook

from bahisync_engine import ai_suggest
from bahisync_engine.export_excel import build_excel
from bahisync_engine.export_tally import build_ledger_masters_xml, build_vouchers_xml
from bahisync_engine.ledgers import SUSPENSE, ClientRules, apply_user_choice, counterparty, suggest
from bahisync_engine.models import Transaction
from bahisync_engine.parser import parse_statement


def txn(narr, debit="0", credit="0", n=1):
    return Transaction(n, date(2026, 4, 1), narr, Decimal(debit), Decimal(credit), Decimal("1000"))


def test_counterparty_extraction():
    assert counterparty("UPI/DR/412345678901/RAMESH KUMAR/SBIN/ramesh@oksbi/Payment") == "Ramesh Kumar"
    assert counterparty("NEFT CR-HDFC0000123-ACME SOLUTIONS PVT LTD-SALARY OCT 2026") == "Acme Solutions PVT LTD"
    assert counterparty("IMPS/P2A/700123/CUSTOMER PAYMENT/KKBK0000958/ANAND ENTERPRISES") == "Anand Enterprises"
    assert counterparty("CHQ PAID 000123 VENKATESWARA HARDWARE") == "Venkateswara Hardware"
    assert counterparty("SMS CHARGES QTR ENDED MAR") == ""


def test_default_rules_and_voucher_types():
    rows = [
        txn("ATM WDL/HYDERABAD AMEERPET", debit="10000.00", n=1),
        txn("GST PAYMENT CBIC CPIN 2604", debit="18000.00", n=2),
        txn("INT.PD:01-07-2026 TO 30-09-2026", credit="1184.00", n=3),
        txn("SOMETHING UNKNOWN 12345", debit="50.00", n=4),
    ]
    groups = suggest(rows)
    assert [r.ledger for r in rows] == ["Cash", "GST Payable", "Bank Interest Received", SUSPENSE]
    assert [r.voucher_type for r in rows] == ["Contra", "Payment", "Receipt", "Payment"]
    assert groups["GST Payable"] == "Duties & Taxes"
    assert "suspense" in rows[3].flags and "needs_review" in rows[3].flags
    assert "needs_review" not in rows[0].flags


def test_client_rules_learn_from_corrections():
    client = ClientRules()
    first = txn("UPI/DR/1111/RAMESH KUMAR/SBIN/r@oksbi/x", debit="500.00")
    suggest([first], client)
    assert first.ledger_source == "party_guess"
    apply_user_choice(first, "Ramesh Kumar (Carpenter)", client, "Sundry Creditors")

    # Next month, same person, different reference number.
    nxt = txn("UPI/DR/2222/RAMESH KUMAR/SBIN/r@oksbi/y", debit="750.00")
    suggest([nxt], ClientRules.from_json(client.to_json()))
    assert nxt.ledger == "Ramesh Kumar (Carpenter)"
    assert nxt.ledger_source == "client_rule"
    assert "needs_review" not in nxt.flags


def test_tally_vouchers_balance_and_signs(fixtures):
    r = parse_statement(fixtures["sbi"].pdf)
    groups = suggest(r.transactions)
    xml = build_vouchers_xml(r.transactions, "Demo Traders", "SBI Current A/c")
    root = ET.fromstring(xml)
    vouchers = root.findall(".//VOUCHER")
    assert len(vouchers) == r.rows_total
    for v in vouchers:
        entries = v.findall("ALLLEDGERENTRIES.LIST")
        amounts = [Decimal(e.findtext("AMOUNT")) for e in entries]
        assert sum(amounts) == 0
        for e in entries:
            positive = e.findtext("ISDEEMEDPOSITIVE") == "Yes"
            assert (Decimal(e.findtext("AMOUNT")) < 0) == positive
    assert root.findtext(".//SVCURRENTCOMPANY") == "Demo Traders"

    masters = ET.fromstring(build_ledger_masters_xml("Demo Traders", "SBI Current A/c", groups, {t.ledger for t in r.transactions}))
    parents = {led.get("NAME"): led.findtext("PARENT") for led in masters.findall(".//LEDGER")}
    assert parents["SBI Current A/c"] == "Bank Accounts"
    assert parents["Cash"] == "Cash-in-Hand"


def test_tally_xml_escapes_special_characters():
    t = txn("PAYMENT TO A & B <TRADERS>", debit="10.00")
    t.ledger = "A & B Traders"
    xml = build_vouchers_xml([t], "Ravi & Co", "Bank")
    root = ET.fromstring(xml)  # must still be valid XML
    assert root.findtext(".//LEDGERNAME") == "A & B Traders"


def test_excel_export(fixtures):
    r = parse_statement(fixtures["kotak"].pdf)
    suggest(r.transactions)
    wb = load_workbook(filename=__import__("io").BytesIO(build_excel(r)))
    ws = wb["Transactions"]
    assert ws.max_row == r.rows_total + 1
    assert ws["H2"].value == "OK"


def test_ai_suggestions_capped_and_flagged(monkeypatch):
    monkeypatch.setenv("ANTHROPIC_API_KEY", "test-key")
    rows = [txn("MISC TRANSFER 99881", debit="200.00", n=7)]
    suggest(rows)

    def handler(request):
        assert request.headers["x-api-key"] == "test-key"
        return httpx.Response(200, json={"content": [{"type": "text", "text": '[{"id": 7, "ledger": "Office Expenses", "group": "Indirect Expenses", "confidence": 0.95}]'}]})

    client = httpx.Client(transport=httpx.MockTransport(handler))
    groups = ai_suggest.suggest_with_ai(rows, ["Office Expenses"], client=client)
    assert rows[0].ledger == "Office Expenses"
    assert rows[0].confidence == ai_suggest.MAX_AI_CONFIDENCE
    assert "needs_review" in rows[0].flags
    assert groups == {"Office Expenses": "Indirect Expenses"}


def test_ai_disabled_without_key(monkeypatch):
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    rows = [txn("MISC", debit="1.00")]
    suggest(rows)
    assert ai_suggest.suggest_with_ai(rows, []) == {}
    assert rows[0].ledger == SUSPENSE
